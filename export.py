from __future__ import annotations

import re
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from proposal_schema import PROPOSAL_SECTIONS, Proposal

COMPANY_BRAND_PLACEHOLDER = "[YOUR COMPANY NAME]"


def build_proposal_document(
    proposal: Proposal,
    client_name: str,
    company_name: str = COMPANY_BRAND_PLACEHOLDER,
) -> Document:
    if not isinstance(proposal, Proposal):
        raise TypeError("proposal must be a validated Proposal instance")
    if not client_name.strip():
        raise ValueError("client_name must not be empty")

    document = Document()
    _configure_document(document)
    _add_header_and_footer(document, company_name, client_name)
    _add_cover_page(document, company_name, client_name)
    document.add_page_break()

    for index, (field, label) in enumerate(PROPOSAL_SECTIONS):
        if index:
            document.add_page_break()
        heading = document.add_heading(label, level=1)
        heading.paragraph_format.keep_with_next = True
        _add_section_body(document, getattr(proposal, field))

    return document


def export_proposal_docx(
    proposal: Proposal,
    client_name: str,
    company_name: str = COMPANY_BRAND_PLACEHOLDER,
) -> bytes:
    document = build_proposal_document(proposal, client_name, company_name)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Aptos"
    normal_style.font.size = Pt(10.5)
    normal_style.paragraph_format.space_after = Pt(8)
    normal_style.paragraph_format.line_spacing = 1.12

    title_style = document.styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style.font.size = Pt(30)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(15, 23, 42)

    heading_style = document.styles["Heading 1"]
    heading_style.font.name = "Aptos Display"
    heading_style.font.size = Pt(18)
    heading_style.font.bold = True
    heading_style.font.color.rgb = RGBColor(37, 99, 235)
    heading_style.paragraph_format.space_before = Pt(0)
    heading_style.paragraph_format.space_after = Pt(12)


def _add_header_and_footer(
    document: Document,
    company_name: str,
    client_name: str,
) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.text = company_name.strip() or COMPANY_BRAND_PLACEHOLDER
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.runs[0]
    header_run.bold = True
    header_run.font.size = Pt(9)
    header_run.font.color.rgb = RGBColor(71, 85, 105)

    footer = section.footer.paragraphs[0]
    footer.text = f"Confidential  •  Prepared for {client_name.strip()}"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.runs[0]
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(100, 116, 139)


def _add_cover_page(
    document: Document,
    company_name: str,
    client_name: str,
) -> None:
    brand = document.add_paragraph(company_name.strip() or COMPANY_BRAND_PLACEHOLDER)
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    brand_run = brand.runs[0]
    brand_run.bold = True
    brand_run.font.size = Pt(14)
    brand_run.font.color.rgb = RGBColor(37, 99, 235)

    document.add_paragraph()
    document.add_paragraph()
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Enterprise Proposal")

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(f"Prepared for {client_name.strip()}")
    subtitle_run.font.size = Pt(16)
    subtitle_run.font.color.rgb = RGBColor(71, 85, 105)

    document.add_paragraph()
    descriptor = document.add_paragraph("Proposal and pilot engagement overview")
    descriptor.alignment = WD_ALIGN_PARAGRAPH.CENTER
    descriptor.runs[0].font.size = Pt(11)
    descriptor.runs[0].font.italic = True


def _add_section_body(document: Document, body: str) -> None:
    blocks = re.split(r"\n\s*\n", body.strip())
    for block in blocks:
        if not block:
            continue
        lines = block.splitlines()
        if lines and all(re.match(r"^\s*(?:[-*•])\s+", line) for line in lines):
            for line in lines:
                item_text = re.sub(r"^\s*(?:[-*•])\s+", "", line)
                document.add_paragraph(item_text, style="List Bullet")
        else:
            paragraph = document.add_paragraph()
            paragraph.add_run("\n".join(lines))
