import unittest
from io import BytesIO

from docx import Document

from export import COMPANY_BRAND_PLACEHOLDER, export_proposal_docx
from proposal_schema import PROPOSAL_SECTIONS, Proposal


class ProposalExportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.proposal = Proposal(
            executive_summary="The client needs a controlled first release.",
            proposed_solution="- Confirm the data boundary\n- Ship a measured pilot",
            pilot_roadmap="Week 1: discovery\nWeek 4: readout",
            why_our_company="The delivery team has shipped governed systems.",
            next_steps="Approve the pilot brief and schedule kickoff.",
        )

    def test_generates_a_readable_docx_from_all_structured_sections(self) -> None:
        document = export_proposal_docx(self.proposal, "Apex Logistics Global")

        self.assertTrue(document.startswith(b"PK"))
        parsed = Document(BytesIO(document))
        headings = [
            paragraph.text
            for paragraph in parsed.paragraphs
            if paragraph.style.name == "Heading 1"
        ]
        self.assertEqual(headings, [label for _, label in PROPOSAL_SECTIONS])
        self.assertIn("Confirm the data boundary", [p.text for p in parsed.paragraphs])
        self.assertIn("Approve the pilot brief and schedule kickoff.", [p.text for p in parsed.paragraphs])

    def test_includes_branding_placeholder_header_and_page_breaks(self) -> None:
        document = export_proposal_docx(self.proposal, "Apex Logistics Global")
        parsed = Document(BytesIO(document))
        header_text = parsed.sections[0].header.paragraphs[0].text
        page_breaks = parsed.element.body.xml.count('w:type="page"')

        self.assertIn(COMPANY_BRAND_PLACEHOLDER, header_text)
        self.assertGreaterEqual(page_breaks, len(PROPOSAL_SECTIONS))
        self.assertIn("Prepared for Apex Logistics Global", [
            paragraph.text
            for paragraph in parsed.paragraphs
        ])

    def test_uses_bullet_format_for_structured_bullet_blocks(self) -> None:
        document = export_proposal_docx(self.proposal, "Apex Logistics Global")
        parsed = Document(BytesIO(document))

        self.assertTrue(
            any(
                paragraph.style.name == "List Bullet"
                and paragraph.text == "Confirm the data boundary"
                for paragraph in parsed.paragraphs
            )
        )


if __name__ == "__main__":
    unittest.main()
